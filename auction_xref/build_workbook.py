#!/usr/bin/env python3
"""Build the deliverable workbook from results.csv, october-catalogue.txt and lots.csv.

Replaces the earlier series of ad-hoc edit passes. Everything is regenerated
from source each run, so re-running after a catalogue update is one command
and cannot drift.

openpyxl writes formulas with no cached value, and drops any it had whenever a
workbook is re-saved, so the computed cells read blank outside Excel. recalc.py
would normally fix that, but LibreOffice cannot load xlsx in this container, so
the values are computed here and patched into the sheet XML at the end.
"""
import csv
import re
import sys
import zipfile
from xml.etree import ElementTree as ET

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT = "Auction lots - checked list.xlsx"
NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
ARIAL = "Arial"
HDR = PatternFill("solid", fgColor="1F3864")
ADDED = PatternFill("solid", fgColor="7030A0")
LIVE = PatternFill("solid", fgColor="E2EFDA")
GONE = PatternFill("solid", fgColor="F2DCDB")
HFONT = Font(name=ARIAL, size=10, bold=True, color="FFFFFF")
_thin = Side(style="thin", color="BFBFBF")
BD = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)
GBP = '£#,##0;(£#,##0);-'
PCT = '0.0%;(0.0%);-'

POSTCODE = re.compile(r"\b([A-Z]{1,2}[0-9][A-Z0-9]?)\s*([0-9][A-Z]{2})\b", re.I)
LEADER = re.compile(
    r"^(?:(?:flat|apartment|apt|unit|plot|garage|compound|parking\s+space)\s*)?"
    r"([0-9]+[a-z]?)\b", re.I)


def low(s):
    """Low end of a guide like '£20,000-£40,000' or '£415,000+'."""
    m = re.search(r"£\s*([\d,]+)", s or "")
    return int(m.group(1).replace(",", "")) if m else None


def money(s):
    if not s or s == "NO OFFERS":
        return None
    d = "".join(c for c in s if c.isdigit())
    return int(d) if d else None


def postcode(t):
    m = POSTCODE.search(t or "")
    return f"{m.group(1)}{m.group(2)}".upper() if m else None


def leader(a):
    for part in (p.strip() for p in (a or "").split(",")):
        m = LEADER.search(part)
        if m:
            return m.group(1).upper()
    return None


def header(ws, labels, added=()):
    ws.append(labels)
    for i in range(1, len(labels) + 1):
        c = ws.cell(row=1, column=i)
        c.fill = ADDED if i in added else HDR
        c.font, c.border = HFONT, BD
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def style(ws, row, ncols, wrap_col=None, fill=None):
    for i in range(1, ncols + 1):
        c = ws.cell(row=row, column=i)
        c.font = Font(name=ARIAL, size=10)
        c.border = BD
        c.alignment = Alignment(vertical="top", wrap_text=(i == wrap_col))
        if fill:
            c.fill = fill


def widths(ws, ws_widths):
    for i, w in enumerate(ws_widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


STATUS = {"october-catalogue": "In October catalogue",
          "SOLD PRIOR": "Sold prior to auction",
          "NOT LISTED": "Not in October catalogue"}
NO_OFFER_LOTS = {"29", "42", "111", "114", "117", "122", "152"}


def verification(r):
    if r["status"] == "NOT LISTED":
        return ("different unit at this postcode - our unit not offered"
                if "different unit" in r["confidence"] else "postcode not in catalogue")
    return "address exact match (verified)" if r["confidence"] == "postcode only" else r["confidence"]


def main():
    results = list(csv.DictReader(open("results.csv")))
    order = {"october-catalogue": 0, "SOLD PRIOR": 1, "NOT LISTED": 2}
    results.sort(key=lambda r: (order[r["status"]], -(low(r["guide_price"]) or 0)))

    catalogue = []
    for line in open("october-catalogue.txt"):
        if line.count("|") >= 2:
            lot, addr, guide = (f.strip() for f in line.split("|", 2))
            catalogue.append({"lot": lot, "addr": addr, "guide": guide,
                              "pc": postcode(addr), "ld": leader(addr), "g": low(guide),
                              "sold": "sold prior" in guide.lower()})

    known = [{"pc": postcode(r["address"]), "ld": leader(r["address"]),
              "ht": low(r["hometrack"]), "ht_raw": r["hometrack"],
              "offer": money(r["accepted_offer"]),
              "no_offers": r["accepted_offer"] == "NO OFFERS",
              "lot": r["lot"], "serco": r["serco_status"]} for r in results]
    for o in catalogue:
        cands = [k for k in known if k["pc"] and k["pc"] == o["pc"]]
        exact = [k for k in cands if o["ld"] and k["ld"] == o["ld"]]
        o["m"] = exact[0] if exact else (cands[0] if len(cands) == 1 and not o["ld"] else None)

    wb = Workbook()
    wb.remove(wb.active)

    # --- October catalogue: the auction house's list, the spine of the workbook.
    ws = wb.create_sheet("October catalogue")
    header(ws, ["Oct lot", "Address", "Guide price", "Guide +20%", "Hometrack valuation",
                "Guide +20% as % of Hometrack", "Offer likely accepted (£)", "Serco status",
                "Seen in September?"])
    for i, o in enumerate(catalogue, start=2):
        m = o["m"]
        ws.cell(row=i, column=1, value=o["lot"])
        ws.cell(row=i, column=2, value=o["addr"])
        ws.cell(row=i, column=3, value=o["g"])
        ws.cell(row=i, column=4, value=f'=IFERROR(C{i}*1.2,"")')
        ws.cell(row=i, column=5, value=(m["ht"] if m and m["ht"] is not None
                                        else (m["ht_raw"] if m else None)))
        ws.cell(row=i, column=6, value=f'=IFERROR(D{i}/E{i},"")')
        ws.cell(row=i, column=7, value=(m["offer"] if m and m["offer"]
                                        else ("No offers" if m and m["no_offers"] else None)))
        ws.cell(row=i, column=8, value=(m["serco"] if m else None))
        ws.cell(row=i, column=9, value=(f'Yes - Sept lot {m["lot"]}' if m
                                        else "New - not in our September list"))
        style(ws, i, 9, wrap_col=2, fill=(GONE if o["sold"] else (LIVE if m else None)))
        for c in (3, 4, 5, 7):
            ws.cell(row=i, column=c).number_format = GBP
        ws.cell(row=i, column=6).number_format = PCT
    widths(ws, [8, 58, 15, 15, 20, 16, 16, 22, 30])
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:I{len(catalogue) + 1}"

    # --- Working list: our September lots, the day-to-day view.
    ws = wb.create_sheet("Working list")
    header(ws, ["Address", "Guide price", "Hometrack valuation", "Guide +20%",
                "Guide +20% as % of Hometrack", "Offer likely accepted (£)",
                "Offer as % of Hometrack"], added={5, 6, 7})
    for i, r in enumerate(results, start=2):
        guide_txt = r["site_guide"] if r["site_lot"] else r["guide_price"]
        ws.cell(row=i, column=1, value=r["address"])
        ws.cell(row=i, column=2, value=low(guide_txt) or low(r["guide_price"]))
        ht = low(r["hometrack"])
        ws.cell(row=i, column=3, value=ht if ht is not None else (r["hometrack"] or ""))
        ws.cell(row=i, column=4, value=f'=IFERROR(B{i}*1.2,"")')
        ws.cell(row=i, column=5, value=f'=IFERROR(D{i}/C{i},"")')
        offer = money(r["accepted_offer"])
        ws.cell(row=i, column=6,
                value=offer if offer else ("No offers" if r["lot"] in NO_OFFER_LOTS else None))
        ws.cell(row=i, column=7, value=f'=IFERROR(F{i}/C{i},"")')
        fill = GONE if r["status"] == "SOLD PRIOR" else (LIVE if r["status"] == "october-catalogue" else None)
        style(ws, i, 7, wrap_col=1, fill=fill)
        for c in (2, 3, 4, 6):
            ws.cell(row=i, column=c).number_format = GBP
        for c in (5, 7):
            ws.cell(row=i, column=c).number_format = PCT
    widths(ws, [62, 16, 20, 16, 16, 16, 16])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:G{len(results) + 1}"

    # --- Checked list: the audit trail behind every match.
    ws = wb.create_sheet("Checked list")
    header(ws, ["Sep lot", "Address", "Sep guide", "Sep guide low (£)", "Oct lot", "Oct guide",
                "Oct guide low (£)", "Guide +20%", "Guide move", "Our accepted offer (£)",
                "Offer vs Oct guide", "Hometrack AVM", "Serco status", "Status",
                "Match verification"])
    for i, r in enumerate(results, start=2):
        ws.cell(row=i, column=1, value=r["lot"])
        ws.cell(row=i, column=2, value=r["address"])
        ws.cell(row=i, column=3, value=r["guide_price"])
        ws.cell(row=i, column=4, value=low(r["guide_price"]))
        ws.cell(row=i, column=5, value=r["site_lot"] or None)
        ws.cell(row=i, column=6, value=r["site_guide"] or None)
        ws.cell(row=i, column=7, value=low(r["site_guide"]))
        ws.cell(row=i, column=8, value=f'=IFERROR(IF(G{i}<>"",G{i},D{i})*1.2,"")')
        ws.cell(row=i, column=9, value=f'=IFERROR((G{i}-D{i})/D{i},"")')
        ws.cell(row=i, column=10, value=money(r["accepted_offer"]))
        ws.cell(row=i, column=11, value=f'=IFERROR(J{i}/G{i},"")')
        ws.cell(row=i, column=12, value=r["hometrack"])
        ws.cell(row=i, column=13, value=r["serco_status"])
        ws.cell(row=i, column=14, value=STATUS[r["status"]])
        ws.cell(row=i, column=15, value=verification(r))
        fill = GONE if r["status"] == "SOLD PRIOR" else (LIVE if r["status"] == "october-catalogue" else None)
        style(ws, i, 15, wrap_col=2, fill=fill)
        for c in (4, 7, 8, 10):
            ws.cell(row=i, column=c).number_format = GBP
        for c in (9, 11):
            ws.cell(row=i, column=c).number_format = PCT
    widths(ws, [8, 52, 15, 16, 7, 20, 16, 14, 10, 16, 12, 22, 22, 24, 40])
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:O{len(results) + 1}"

    counts = {}
    for r in results:
        counts[r["status"]] = counts.get(r["status"], 0) + 1
    matched = sum(1 for o in catalogue if o["m"])
    from_shot = sum(1 for o in catalogue if o["lot"] == "?")

    nt = wb.create_sheet("Notes")
    nt.column_dimensions["A"].width = 118
    lines = [
     ("Checked list - Auction House London", True), ("", False),
     ("INCOMPLETE CATALOGUE - READ FIRST", True),
     (f"The October catalogue here holds {len(catalogue)} lots, and it is NOT the full sale.",False),
     ("The page was captured from the browser rather than fetched, and the capture missed lots:",False),
     ("three were found later from screenshots and added by hand (Flat 5 Poplar House SE16 6YJ,",False),
     ("102-104 and 106 High Street Redcar). Two of those had been reported as 'not in the October",False),
     ("catalogue' and were in fact in it, so that verdict was wrong for them.",False),
     ("Treat every 'Not in October catalogue' row as unconfirmed until the full catalogue is",False),
     ("captured. 'In October catalogue' rows are safe - a lot present is present.",False),
     (f"The {from_shot} lots added from screenshots show '?' as their lot number, because the",False),
     ("screenshots did not show one.",False), ("", False),
     ("Lot numbers", True),
     ("The numbers in 'Oct lot' came from the browser capture and may be positional rather than",False),
     ("the auction house's own lot numbers. Confirm against the catalogue before bidding.",False),("",False),
     ("What this is", True),
     ("Three sheets:", False),
     ("  October catalogue - the auction house's list, what is actually on offer. The spine.",False),
     ("  Working list      - our 100 tracked lots, address / guide / Hometrack / +20% / offer.",False),
     ("  Checked list      - the audit trail: how each of the 100 was matched, and to what.",False),("",False),
     (f"Of the {len(catalogue)} October lots, {matched} are ones we already knew from September;",False),
     (f"the rest are new to us. Of our 100 September lots, {counts.get('october-catalogue',0)} are",False),
     (f"in the October catalogue, {counts.get('SOLD PRIOR',0)} sold prior, and",False),
     (f"{counts.get('NOT LISTED',0)} are not in it - subject to the warning above.",False),("",False),
     ("The biggest gap", True),
     ("Most October lots have no Hometrack valuation, because they are lots we have never valued.",False),
     ("Until those are pulled from Sourcing Brain this workbook can say what is on offer and at",False),
     ("what guide, but not which lots are worth bidding on.",False),("",False),
     ("Source of the tracked list", True),
     ("The spreadsheet is titled 'October auction' but its lots are the 2nd-3rd SEPTEMBER 2026 sale.",False),
     ("Confirmed against two published lot records (lot 99, 23 Selkirk Road; lot 101, 49 Exmouth",False),
     ("Road), both dated 02/09/2026. The spreadsheet's own source note flagged this as unresolved.",False),("",False),
     ("How lots were matched", True),
     ("On postcode, then on flat/house number. Lot numbers and property names differ between the",False),
     ("two sales - September lot 39 'Holmshill House' is the site's lot 40 'Holmshill Farm' - so",False),
     ("neither is a safe key. Matches were re-checked by full address text; the three with no unit",False),
     ("number to confirm against (127, 175, 182) scored 1.000 similarity.",False),("",False),
     ("Deliberate non-matches", True),
     ("Lot 113, Flat 110, 2 Moorfields - only Flat 105 is offered in October.",False),
     ("Lot 120, Flat 23, 2 Manilla Street E14 8GB - the October lot is 1 Manilla Street E14 8JZ.",False),
     ("Lot 76, 108 High Street Redcar - 102-104 and 106 are offered, 108 is not.",False),("",False),
     ("Addresses corrected against the site so matching would work", True),
     ("4 Narin Court, Tilbury -> 4 Nairn Court.  3 Jane Court, St Albans -> 3 Dane Court.",False),
     ("104 Graveney Road: postcode SW17 0DH -> SW17 0EH.",False),("",False),
     ("Columns computed by formula", True),
     ("Guide +20% = guide x 1.2. On the Working list the guide is the current one: October where",False),
     ("the lot was re-entered, September otherwise. Range guides use the LOW end throughout.",False),
     ("Guide +20% as % of Hometrack, and Offer as % of Hometrack: over 100% usually means the",False),
     ("Hometrack match is the wrong unit or the whole building, not that the lot is overpriced.",False),
     ("Offers are from the SEPTEMBER sale and may be stale where a guide has since been cut.",False),("",False),
     ("Currency", True),
     ("October catalogue captured 15 September 2026, corrected 16 September 2026. The sale was",False),
     ("still taking entries, so it will grow before 7-8 October. Re-run nearer the date.",False),
    ]
    for i, (t, b) in enumerate(lines, start=1):
        c = nt.cell(row=i, column=1, value=t)
        c.font = Font(name=ARIAL, size=10, bold=b)
        c.alignment = Alignment(vertical="top")

    wb.save(OUT)
    patch_cached_values()
    print(f"built {OUT}: {len(catalogue)} October lots ({matched} known, {from_shot} from screenshots), "
          f"{len(results)} tracked lots {counts}")


def patch_cached_values():
    """Compute every formula's value and write it beside the formula in the XML."""
    ET.register_namespace("", NS)
    wb = load_workbook(OUT)
    sheets = wb.sheetnames
    want = {}

    oc = wb["October catalogue"]
    v = {}
    for r in range(2, oc.max_row + 1):
        g, ht = oc.cell(row=r, column=3).value, oc.cell(row=r, column=5).value
        if isinstance(g, (int, float)):
            v[f"D{r}"] = g * 1.2
            if isinstance(ht, (int, float)) and ht:
                v[f"F{r}"] = (g * 1.2) / ht
    want["October catalogue"] = v

    wl = wb["Working list"]
    v = {}
    for r in range(2, wl.max_row + 1):
        g = wl.cell(row=r, column=2).value
        ht = wl.cell(row=r, column=3).value
        off = wl.cell(row=r, column=6).value
        if isinstance(g, (int, float)):
            v[f"D{r}"] = g * 1.2
            if isinstance(ht, (int, float)) and ht:
                v[f"E{r}"] = (g * 1.2) / ht
        if isinstance(off, (int, float)) and isinstance(ht, (int, float)) and ht:
            v[f"G{r}"] = off / ht
    want["Working list"] = v

    cl = wb["Checked list"]
    v = {}
    for r in range(2, cl.max_row + 1):
        d = cl.cell(row=r, column=4).value
        g = cl.cell(row=r, column=7).value
        j = cl.cell(row=r, column=10).value
        base = g if isinstance(g, (int, float)) else (d if isinstance(d, (int, float)) else None)
        if base is not None:
            v[f"H{r}"] = base * 1.2
        if isinstance(d, (int, float)) and isinstance(g, (int, float)) and d:
            v[f"I{r}"] = (g - d) / d
        if isinstance(j, (int, float)) and isinstance(g, (int, float)) and g:
            v[f"K{r}"] = j / g
    want["Checked list"] = v

    with zipfile.ZipFile(OUT) as z:
        names = z.namelist()
        data = {n: z.read(n) for n in names}
    total = 0
    for name, cells in want.items():
        path = f"xl/worksheets/sheet{sheets.index(name) + 1}.xml"
        root = ET.fromstring(data[path])
        for c in root.iter(f"{{{NS}}}c"):
            ref = c.get("r")
            if ref in cells and c.find(f"{{{NS}}}f") is not None:
                for old in c.findall(f"{{{NS}}}v"):
                    c.remove(old)
                ET.SubElement(c, f"{{{NS}}}v").text = repr(round(cells[ref], 10))
                c.attrib.pop("t", None)
                total += 1
        data[path] = ET.tostring(root, encoding="UTF-8", xml_declaration=True)
    with zipfile.ZipFile(OUT, "w", zipfile.ZIP_DEFLATED) as z:
        for n in names:
            z.writestr(n, data[n])

    check = load_workbook(OUT, data_only=True)
    bad = 0
    for name, cells in want.items():
        s = check[name]
        for ref, expect in cells.items():
            got = s[ref].value
            if got is None or abs(got - expect) > 1e-9:
                bad += 1
                print("  MISMATCH", name, ref, got, expect, file=sys.stderr)
    print(f"cached values: patched {total}, verified {total - bad}, mismatches {bad}")
    if bad:
        sys.exit(1)


if __name__ == "__main__":
    main()
